from datetime import datetime, timedelta
from io import BytesIO
import json

from flask import Blueprint, request, send_file
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import DateTime, func, or_

from app.database.database import db
from app.models.audit_event import AuditEvent, serialize_record
from app.models.cita import Cita
from app.models.comprobante import Comprobante
from app.models.detalle_reserva import Detalle_Reserva
from app.models.inventario import Inventario
from app.models.lote import Lote
from app.models.prenda import Prenda
from app.models.prenda_imagen import PrendaImagen
from app.models.reserva import Reserva
from app.models.usuarios import Usuarios
from app.utils.auth_middleware import admin_required
from app.utils.response import response_error

reports_bp = Blueprint('reports', __name__, url_prefix='/api/reportes')

_EXCEL_MAX_DATA_ROWS = 1_048_575
_EXCEL_MIMETYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
_HEADER_FILL = PatternFill(fill_type='solid', fgColor='155D3A')
_ALTERNATE_FILL = PatternFill(fill_type='solid', fgColor='EFF5F1')
_HEADER_FONT = Font(color='FFFFFF', bold=True)
_BODY_ALIGNMENT = Alignment(vertical='top', wrap_text=True)
_HEADER_ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)
_REPORT_MODELS = (
    Reserva,
    Comprobante,
    Cita,
    Inventario,
    Lote,
    Prenda,
    PrendaImagen,
    Usuarios,
)
_BUSINESS_DATE_FIELDS = {
    Reserva: ('fecha_reserva', 'fecha_evento', 'fecha_inicio', 'fecha_fin', 'fecha_devolucion'),
    Cita: ('fecha_cita',),
}


def _excel_safe(value):
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def _display_value(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    return _excel_safe(value)


def _append_styled_row(sheet, values, row_number, header=False):
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=_display_value(value))
        if header:
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = _HEADER_ALIGNMENT
        else:
            cell.alignment = _BODY_ALIGNMENT
            if isinstance(value, datetime):
                cell.number_format = 'yyyy-mm-dd hh:mm:ss'
            if row_number % 2 == 0:
                cell.fill = _ALTERNATE_FILL
        cells.append(cell)
    sheet.append(cells)


def _create_report_sheet(workbook, title, headers, widths):
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A2'
    sheet.row_dimensions[1].height = 32
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    _append_styled_row(sheet, headers, 1, header=True)
    return sheet


def _audit_change_details(item):
    changes = item.changes
    if not isinstance(changes, dict):
        yield 'Detalle', None, changes
        return

    if item.action in ('create', 'delete'):
        key = 'after' if item.action == 'create' else 'before'
        snapshot = changes.get(key, {})
        if isinstance(snapshot, dict):
            for field, value in snapshot.items():
                yield field, None if item.action == 'create' else value, value if item.action == 'create' else None
            return

    for field, values in changes.items():
        if isinstance(values, dict) and ('before' in values or 'after' in values):
            yield field, values.get('before'), values.get('after')
        else:
            yield field, None, values


def _date_bounds(column, start, end):
    if isinstance(column.type, DateTime):
        return start, end
    return start.date(), end.date()


def _matching_date_labels(instance, fields, start, end):
    labels = []
    for field in fields:
        value = getattr(instance, field, None)
        if value is None:
            continue
        lower_bound, upper_bound = (start, end) if isinstance(value, datetime) else (start.date(), end.date())
        if lower_bound <= value < upper_bound:
            formatted_value = value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat()
            labels.append(f'{field}={formatted_value}')
    return '; '.join(labels)


def _current_period_records(start, end):
    for model in _REPORT_MODELS:
        date_fields = [
            field
            for field in ('created_at', 'updated_at', *_BUSINESS_DATE_FIELDS.get(model, ()))
            if hasattr(model, field)
        ]
        date_columns = [getattr(model, field) for field in date_fields]
        if not date_columns:
            continue

        date_filters = []
        for column in date_columns:
            lower_bound, upper_bound = _date_bounds(column, start, end)
            date_filters.append((column >= lower_bound) & (column < upper_bound))
        query = model.query.filter(or_(*date_filters)).order_by(*model.__mapper__.primary_key)

        for instance in query.yield_per(1000):
            yield instance, _matching_date_labels(instance, date_fields, start, end)

    reservation_date_fields = ('created_at', 'updated_at', *_BUSINESS_DATE_FIELDS[Reserva])
    reservation_filters = []
    for field in reservation_date_fields:
        column = getattr(Reserva, field)
        lower_bound, upper_bound = _date_bounds(column, start, end)
        reservation_filters.append((column >= lower_bound) & (column < upper_bound))
    details = (
        db.session.query(Detalle_Reserva, Reserva)
        .join(Reserva, Detalle_Reserva.idReserva == Reserva.idReserva)
        .filter(or_(*reservation_filters))
        .order_by(*Detalle_Reserva.__mapper__.primary_key)
        .yield_per(1000)
    )
    for detail, reservation in details:
        yield detail, f'Reserva: {_matching_date_labels(reservation, reservation_date_fields, start, end)}'


@reports_bp.route('/trimestral/excel', methods=['GET'])
@admin_required
def download_quarterly_report():
    year = request.args.get('year', type=int)
    quarter = request.args.get('quarter', type=int)
    if year is None or quarter is None:
        return response_error("Los parámetros 'year' y 'quarter' son requeridos", 400)
    if year < 1 or year > 9998:
        return response_error("'year' debe estar entre 1 y 9998", 400)
    if quarter not in (1, 2, 3, 4):
        return response_error("'quarter' debe estar entre 1 y 4", 400)

    start_month = (quarter - 1) * 3 + 1
    start = datetime(year, start_month, 1)
    if quarter == 4:
        end = datetime(year + 1, 1, 1)
    else:
        end = datetime(year, start_month + 3, 1)
    period = (AuditEvent.occurred_at >= start, AuditEvent.occurred_at < end)

    event_count = AuditEvent.query.filter(*period).count()
    if event_count > _EXCEL_MAX_DATA_ROWS:
        return response_error(
            "El trimestre supera el máximo de filas que admite una hoja de Excel",
            413,
        )

    workbook = Workbook(write_only=True)
    summary_sheet = workbook.create_sheet('Resumen')
    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.freeze_panes = 'A7'
    summary_sheet.column_dimensions['A'].width = 36
    summary_sheet.column_dimensions['B'].width = 64
    summary_rows = [
        ['Reporte trimestral', f'T{quarter} {year}'],
        ['Período UTC', f'{start:%Y-%m-%d} a {(end - timedelta(days=1)):%Y-%m-%d}'],
        ['Total de eventos registrados', event_count],
        [
            'Nota',
            'El historial detallado solo existe desde la instalación; los registros actuales usan fechas disponibles.',
        ],
        [],
        ['Tipo de registro', 'Acción', 'Cantidad'],
    ]
    for row_number, row in enumerate(summary_rows, start=1):
        _append_styled_row(summary_sheet, row, row_number, header=(row_number == 1 or row_number == 6))
    summary_sheet.row_dimensions[1].height = 26
    summary_sheet.row_dimensions[4].height = 36

    grouped_counts = (
        db.session.query(AuditEvent.entity_type, AuditEvent.action, func.count(AuditEvent.idAuditEvent))
        .filter(*period)
        .group_by(AuditEvent.entity_type, AuditEvent.action)
        .order_by(AuditEvent.entity_type, AuditEvent.action)
        .all()
    )
    for row_number, (entity_type, action, count) in enumerate(grouped_counts, start=7):
        _append_styled_row(summary_sheet, [entity_type, action, count], row_number)
    summary_sheet.auto_filter.ref = f'A6:C{max(6, 6 + len(grouped_counts))}'

    history_sheet = _create_report_sheet(workbook, 'Historial', [
        'Fecha y hora UTC',
        'Tipo de registro',
        'ID del registro',
        'Acción',
        'ID del usuario responsable',
    ], [22, 24, 18, 16, 24])
    history_count = 0

    events = (
        AuditEvent.query.filter(*period)
        .order_by(AuditEvent.occurred_at, AuditEvent.idAuditEvent)
        .yield_per(1000)
    )
    detail_sheet = _create_report_sheet(workbook, 'Detalle de cambios', [
        'Fecha y hora UTC',
        'Tipo de registro',
        'ID del registro',
        'Acción',
        'Campo',
        'Valor anterior',
        'Valor nuevo',
    ], [22, 24, 18, 16, 32, 52, 52])
    detail_count = 0
    for item in events:
        history_count += 1
        _append_styled_row(history_sheet, [
            item.occurred_at,
            item.entity_type,
            item.entity_id,
            item.action,
            item.actor_user_id,
        ], history_count + 1)
        for field, before, after in _audit_change_details(item):
            detail_count += 1
            if detail_count > _EXCEL_MAX_DATA_ROWS:
                return response_error(
                    "El trimestre supera el máximo de filas que admite una hoja de Excel",
                    413,
                )
            _append_styled_row(detail_sheet, [
                item.occurred_at,
                item.entity_type,
                item.entity_id,
                item.action,
                field,
                before,
                after,
            ], detail_count + 1)
    history_sheet.auto_filter.ref = f'A1:E{event_count + 1}'
    detail_sheet.auto_filter.ref = f'A1:G{detail_count + 1}'

    current_records_sheet = _create_report_sheet(workbook, 'Registros actuales', [
        'Tipo de registro',
        'ID del registro',
        'Fechas coincidentes UTC',
    ], [24, 20, 72])
    current_data_sheet = _create_report_sheet(workbook, 'Datos actuales', [
        'Tipo de registro',
        'ID del registro',
        'Campo',
        'Valor actual',
    ], [24, 20, 36, 64])
    current_record_count = 0
    current_field_count = 0
    for instance, reference_at in _current_period_records(start, end):
        current_record_count += 1
        if current_record_count > _EXCEL_MAX_DATA_ROWS:
            return response_error(
                "El trimestre supera el máximo de filas que admite una hoja de Excel",
                413,
            )
        identity = type(instance).__mapper__.primary_key
        record_id = ':'.join(str(getattr(instance, column.key)) for column in identity)
        entity_type = type(instance).__name__
        _append_styled_row(current_records_sheet, [
            entity_type,
            record_id,
            reference_at,
        ], current_record_count + 1)
        for field, value in serialize_record(instance).items():
            current_field_count += 1
            if current_field_count > _EXCEL_MAX_DATA_ROWS:
                return response_error(
                    "El trimestre supera el máximo de filas que admite una hoja de Excel",
                    413,
                )
            _append_styled_row(current_data_sheet, [
                entity_type,
                record_id,
                field,
                value,
            ], current_field_count + 1)
    current_records_sheet.auto_filter.ref = f'A1:C{current_record_count + 1}'
    current_data_sheet.auto_filter.ref = f'A1:D{current_field_count + 1}'

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype=_EXCEL_MIMETYPE,
        as_attachment=True,
        download_name=f'reporte_trimestral_{year}_T{quarter}.xlsx',
    )
